"""Build the editable Hanui Relay hackathon decks from local synthetic assets."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
W, H = 13.333, 7.5
GREEN = RGBColor(19, 83, 61)
MINT = RGBColor(224, 240, 232)
INK = RGBColor(28, 42, 36)
MUTED = RGBColor(91, 110, 101)
LINE = RGBColor(218, 227, 221)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 248)
FONT = "맑은 고딕"

def box(slide, x, y, w, h, fill=WHITE, radius=False, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background() if line is None else None
    if line is not None: sh.line.color.rgb = line
    if radius:
        try: sh.adjustments[0] = 0.12
        except Exception: pass
    return sh

def text(slide, s, x, y, w, h, size=24, bold=False, color=INK, align=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.01)
    for i, row in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = row; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = color
        p.space_after = Pt(10)
        if align is not None: p.alignment = align
    return tb

def base(prs, section, title, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, W, .13, GREEN)
    text(slide, "HANUI RELAY  /  " + section.upper(), .55, .39, 8, .35, 11, True, GREEN)
    text(slide, title, .55, .92, 12.1, .83, 31, True)
    box(slide, .55, 7.09, 12.2, .012, LINE)
    text(slide, "DevDay Exchange Seoul · Track 2 / AI for Human Life", .55, 7.14, 9, .25, 9, False, MUTED)
    text(slide, f"{number:02d} / {total:02d}", 11.55, 7.14, 1.2, .25, 9, False, MUTED, PP_ALIGN.RIGHT)
    return slide

def pill(slide, s, x, y, w):
    box(slide, x, y, w, .43, MINT, True)
    text(slide, s, x+.13, y+.06, w-.25, .3, 13, True, GREEN)

def card(slide, head, body, x, y, w, h, accent=False):
    box(slide, x, y, w, h, WHITE, True, LINE)
    if accent: box(slide, x, y, .07, h, GREEN)
    text(slide, head, x+.25, y+.25, w-.5, .48, 20, True, GREEN)
    text(slide, body, x+.25, y+.87, w-.5, h-1.04, 17, False, INK)

def shot(slide, name, x, y, w, h, crop_top=.0, crop_bottom=.0):
    path = ASSETS / name
    # Cropping into a separate local image keeps the embedded screenshot crisp.
    im = Image.open(path).convert("RGB")
    a=int(im.height*crop_top); b=int(im.height*(1-crop_bottom))
    im=im.crop((0,a,im.width,b))
    target=ASSETS/(Path(name).stem+f"-crop-{a}-{b}.png")
    if not target.exists(): im.save(target)
    box(slide, x-.04, y-.04, w+.08, h+.08, WHITE, True, LINE)
    ratio=min(w/(im.width/im.height),h)
    iw=ratio*im.width/im.height
    slide.shapes.add_picture(str(target), Inches(x+(w-iw)/2), Inches(y), width=Inches(iw), height=Inches(ratio))

def points(slide, rows, x=.75, y=2.05, w=6.0, size=21, gap=.89):
    for i,s in enumerate(rows):
        box(slide,x,y+i*gap+.09,.12,.12,GREEN,True)
        text(slide,s,x+.3,y+i*gap,w-.3,.7,size)

def make_4():
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    n=8
    s=base(prs,"문제 → 제품","병원 밖 14일을, 다음 진료에 연결하다",1,n)
    text(s,"Hanui Relay",.75,2.15,7.4,.92,44,True,GREEN)
    text(s,"한의사 진단·지침 → AI 확인 질문 → 환자 답변 → 내원 리포트",.78,3.25,11.8,1.2,25)
    pill(s,"대상: 생활 지침을 받은 환자와 다음 방문의 한의사",.78,5.3,7.25)
    text(s,"합성 데모 / 의료 진단·처방 아님",.78,6.08,7,.35,14,False,MUTED)
    s=base(prs,"문제 정의","방문 사이의 생활 기록이 진료 맥락에서 끊긴다",2,n)
    card(s,"환자","지침을 들었지만 며칠 뒤 무엇을 했는지 기억하기 어렵습니다.",.65,2.2,5.8,2.6,True)
    card(s,"한의사","다음 내원 때 날짜별 자기보고와 빠진 날을 함께 보기 어렵습니다.",6.85,2.2,5.8,2.6,True)
    text(s,"제품 가설: 생활 지침과 환자 발언을 연결하면 방문 준비의 부담을 줄일 수 있다.",.8,5.3,11.8,.95,22,True,GREEN)
    s=base(prs,"해결 방식","한의학의 생활 지침을 대화와 근거에 연결한다",3,n)
    points(s,["한의사 진단·평가와 생활지침 입력","Hanui가 지침에 맞춘 첫 확인 질문","환자 답변 → 기록·짧은 코칭·고문헌 탐색","서버가 답변 원문·날짜·수치·인용 위치 검증"],.85,2.0,11.5,22,.98)
    pill(s,"선택 이유: 생활 지침 + 한자 원전 장벽 + 환자·한의사 접점",.85,6.1,9.4)
    pill(s,"DB 99건",10.45,6.1,2.0)
    s=base(prs,"데모 ①","한의사 지침을 받은 Hanui가 먼저 묻는다",4,n)
    shot(s,"opening-sequence-v2.png",6.3,2.16,6.25,3.85)
    points(s,["합성 한의사 진단: 기능성 소화불량","Hanui: ‘어젯밤 몇 시에 주무셨어요?’","환자: ‘새벽 1시·5시간’ → 다음 질문"],.75,2.32,5.25,19,1.12)
    text(s,"합성 의료진 입력·환자 대화 / AI가 새 진단을 내리지 않음",.8,6.55,11.6,.36,12,False,MUTED)
    s=base(prs,"데모 ②","14일 리포트의 각 문장은 환자 원문으로 돌아간다",5,n)
    shot(s,"report-evidence-detail.png",6.4,2.04,6.05,4.45)
    text(s,"6 / 8",.8,2.16,4.8,.85,38,True,GREEN)
    text(s,"기록일 중 지침 준수 75%",.8,3.08,5.1,.55,21,True)
    text(s,"기록 없는 6일은 실패로 세지 않음",.8,3.92,5.3,.85,19)
    text(s,"원문 · 날짜 · 문자 위치 확인",.8,5.25,5.1,.8,19)
    text(s,"합성 14일 예시 / 효과 입증 수치 아님",.8,6.45,5.5,.3,12,False,MUTED)
    s=base(prs,"데모 ③","다음 방문을 준비하고 개인 달력으로 이어간다",6,n)
    shot(s,"showcase-calendar.png",6.25,1.9,6.23,4.83,0,.33)
    points(s,["실제 공개 병원 정보 검색","전화·예약 링크와 예약 준비","확정은 사용자가 보고한 뒤 기록","일정 추가·수정·삭제, ICS 내보내기"],.75,2.05,5.2,19,.94)
    text(s,"병원에 실제 접수하거나 길찾기를 제공하지 않음",.8,6.46,11.5,.3,12,False,MUTED)
    s=base(prs,"검증과 Codex","실제 동작을 확인하고, 실패 경계를 명시했다",7,n)
    card(s,"실행 증거","단위 테스트 76개 통과\n실제 Luna 선행질문→코칭 1회\n지침 4턴·달력 8턴·고문헌 1턴",.7,2.0,5.78,3.4,True)
    card(s,"Codex 활용·개선","구현 → 원문·테스트 대조\n사용자 가독성 피드백 반영: 대화창 확대·보조정보 접기",6.77,2.0,5.78,3.4,True)
    text(s,"실패 예: 같은 날 상충 발언 → 실천 판정 보류. 임상 해석 검증은 다음 과제.",.8,5.73,11.8,.8,18,True,GREEN)
    s=base(prs,"도입 가설","환자와 한의사의 다음 만남을 더 준비된 대화로",8,n)
    points(s,["환자: 내 기록을 이해하고 방문 전 정리","잠재 도입 주체 한의원: 내원 전 기록 확인·반복 질문 부담 감소 가설","파일럿: 확인 시간·재사용률·모델 비용 측정"],.82,2.1,11.8,21,1.1)
    pill(s,"재방문·지침 변경마다 누적 기록을 다시 연결",.82,5.9,8.0)
    prs.save(ROOT/"Hanui-Relay-4min.pptx")

def make_10():
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H); n=14
    s=base(prs,"제안","병원 밖 14일을, 다음 진료에 연결하다",1,n)
    text(s,"Hanui Relay",.8,2.17,9.3,.95,45,True,GREEN)
    text(s,"한의사 진단·지침에서 첫 질문, 답변, 내원 리포트까지",.82,3.35,11.7,.68,26)
    pill(s,"Track 2 · AI for Human Life",.82,5.34,5.25)
    s=base(prs,"사용자 문제","방문 사이의 기록은 쉽게 사라진다",2,n)
    card(s,"환자","생활 지침을 기억하고 실천했는지 날짜별로 설명해야 합니다.",.67,2.12,5.85,3.15,True)
    card(s,"한의사","다음 방문에서 자기보고의 맥락과 빈 기간을 짧은 시간에 파악해야 합니다.",6.78,2.12,5.85,3.15,True)
    text(s,"이는 현장 문제에 대한 제품 가설이며, 실제 사용자 조사 결과로 포장하지 않습니다.",.8,5.8,11.7,.75,17,False,MUTED)
    s=base(prs,"왜 한의학인가","생활 지침과 한자 원전의 정보 장벽이 함께 있다",3,n)
    points(s,["수면·식사·활동 등 일상에서 이어지는 지침","고문헌 원문을 환자가 바로 읽기 어려움","환자 설명과 한의사 확인 사이에 출처가 필요"],.85,2.1,11.6,23,1.1)
    pill(s,"목표: 해석과 기록 접근성 / 진단·처방·효능 주장 아님",.84,5.84,9.2)
    s=base(prs,"사용 흐름","한 번의 지침이 다음 방문 자료가 된다",4,n)
    labels=[("1  한의사 입력","진단·평가·지침 저장"),("2  첫 질문","Hanui가 먼저 묻고 환자가 답변"),("3  기록·근거","실천 코칭·고문헌 탐색"),("4  내원","14일 리포트·방문 준비")]
    for i,(h,b) in enumerate(labels): card(s,h,b,.65+i*3.16,2.27,2.9,2.78,i==3)
    text(s,"대화·목표·기록·일정은 세션별 SQLite에 저장되고 재접속 뒤에도 복원됩니다.",.8,5.55,11.7,.65,18)
    s=base(prs,"AI 판단","Luna High가 고르고, 서버가 원문을 검증한다",5,n)
    card(s,"Luna High","한의사 진단·지침 기반 첫 질문\n답변 기록 후보·DB 검색어 선택\n공개 병원 정보 웹 검색",.68,2.04,5.82,3.7,True)
    card(s,"서버 검증","환자 원문·날짜·수치 재확인\nDB 본문 인용 문자열·위치 확인\n세션 소유권·일정 상태 검증",6.78,2.04,5.82,3.7,True)
    text(s,"공식 Codex CLI / 현재 ChatGPT 계정 / gpt-6-luna · high · OpenAI",.8,6.14,11.6,.45,15,True,GREEN)
    s=base(prs,"자료와 출처","99건을 전문 검색하고, 출처의 범위를 드러낸다",6,n)
    text(s,"42",.78,2.08,2.3,1.0,45,True,GREEN); text(s,"현대 자료 자체 요약",.8,3.2,4.7,.5,21)
    text(s,"57",6.42,2.08,2.3,1.0,45,True,GREEN); text(s,"고문헌 원문 발췌",6.44,3.2,4.7,.5,21)
    box(s,.78,4.1,11.6,1.4,WHITE,True,LINE)
    text(s,"동의보감 · 황제내경 · 신농본초경",1.08,4.42,10.9,.47,22,True)
    text(s,"한국어 해석은 AI 설명이고, 원문 일치는 문자·위치 수준의 확인입니다.",.8,5.88,11.6,.75,17,False,MUTED)
    s=base(prs,"데모 ①","한의사 진단·지침에서 첫 확인 질문이 시작된다",7,n)
    shot(s,"opening-sequence-v2.png",6.15,2.04,6.3,3.9)
    points(s,["한의사 진단·23시 전 취침 지침","Hanui: ‘어젯밤 몇 시에 주무셨어요?’","‘새벽 1시·5시간’ → 기록·다음 질문"],.75,2.33,5.3,19,1.09)
    text(s,"합성 환자·지침 예시",.78,6.53,5.1,.3,12,False,MUTED)
    s=base(prs,"데모 ②","대화 중 고문헌을 찾아 원문 위치까지 확인",8,n)
    shot(s,"chat-source-detail.png",6.2,2.25,6.24,3.55)
    points(s,["Luna가 질문의 검색어 선택","SQLite 원문·독해 전문 검색","원문 인용·장절·문자 위치 표시"],.75,2.33,5.2,19,1.08)
    text(s,"역사 기록과 현대 임상 근거를 구분",.78,6.48,5.5,.36,13,False,MUTED)
    s=base(prs,"데모 ③","14일 리포트는 집계부터 문장까지 추적 가능",9,n)
    shot(s,"report-evidence-detail.png",6.22,2.0,6.22,4.52)
    text(s,"75%",.79,2.14,4.65,.9,43,True,GREEN)
    text(s,"8일 기록 중 6일 준수",.81,3.16,5.1,.53,21,True)
    text(s,"6일 기록 없음 · 상충 기록 보류",.81,4.1,5.2,.74,19)
    text(s,"각 근거에서 원래 환자 대화로 이동",.81,5.26,5.18,.72,19)
    text(s,"합성 예시 / 치료 효과 지표 아님",.81,6.49,5.1,.32,12,False,MUTED)
    s=base(prs,"안전한 판정","빈 날과 서로 다른 진술을 임의로 채우지 않는다",10,n)
    shot(s,"conflict-detail.png",6.23,2.02,6.2,4.5)
    points(s,["기록 없음 → 실천율 분모 제외","같은 날 상충 수치 → 보류","가정·미래·타인의 말 → 미기록","사용자가 잘못 잡힌 관찰을 제외"],.76,2.1,5.25,19,.91)
    text(s,"의료 판단 정확성·실사용 효과는 앞으로 검증",.78,6.48,5.4,.38,12,False,MUTED)
    s=base(prs,"방문 준비","병원 정보에서 자체 캘린더와 ICS까지",11,n)
    shot(s,"showcase-calendar.png",6.25,1.89,6.18,4.84,0,.33)
    points(s,["Luna High의 실제 공개 정보 검색","전화·예약 링크로 외부 접수","사용자 확정 보고 후 방문 일정","대화로 일정 추가·수정·삭제"],.75,2.1,5.2,19,.91)
    text(s,"예약 준비와 병원의 실제 접수는 구분",.78,6.48,5.7,.33,12,False,MUTED)
    s=base(prs,"검증과 개선","작동한 범위와 남은 가설을 분리했다",12,n)
    card(s,"확인","단위 테스트 76개 통과\n실제 Luna 선행질문→코칭 1회\n지침 4턴·달력 8턴·고문헌 1턴",.68,2.01,5.8,3.65,True)
    card(s,"수정·남은 과제","상충 기록 보류·빈 날 명시\n인용 원문 일치·세션 분리 확인\n임상 해석·사용자 효과는 미검증",6.77,2.01,5.8,3.65,True)
    text(s,"실제 계정 연결 검증과 화면용 합성 데이터 검증은 별도로 수행했습니다.",.8,6.06,11.6,.52,16,False,MUTED)
    s=base(prs,"Codex 과정","Codex가 구현하고, 원문과 작동을 대조했다",13,n)
    points(s,["Codex: 서버·DB·UI·검색·테스트·자료 정리","Codex 검토: 원전 판본·인용 위치·생활 판정·화면 동작","사용자 피드백: 핵심 집중, 대화창 확대, 해석 우선 표시"],.82,2.14,11.55,21,1.09)
    pill(s,"당일 코드·검증과 사전 기획을 구분해 기록",.82,5.83,8.4)
    s=base(prs,"도입 가능성","다음 만남의 준비 시간을 줄일 수 있는지 검증한다",14,n)
    points(s,["잠재 도입·비용 주체: 한의원 / 사용자: 환자","가치 가설: 내원 전 기록 확인과 반복 질문 부담 감소","작은 파일럿: 확인 시간·재사용률·모델 호출 비용 측정"],.82,2.1,11.55,21,1.07)
    pill(s,"지침 변경·재방문마다 기록 연결 · 가격과 효과는 아직 미측정",.82,5.9,10.65)
    prs.save(ROOT/"Hanui-Relay-10min.pptx")

def add_speaker_notes(deck_name, script_name):
    prs = Presentation(ROOT/deck_name)
    rows = [line for line in (ROOT/script_name).read_text(encoding="utf-8").splitlines()
            if line.startswith("| ") and "초)" in line]
    for slide, row in zip(prs.slides, rows):
        parts = [p.strip() for p in row.strip().strip("|").split("|")]
        slide.notes_slide.notes_text_frame.text = f"권장 시간: {parts[0]}\n\n{parts[2]}"
    prs.save(ROOT/deck_name)

if __name__ == "__main__":
    make_4(); make_10()
    add_speaker_notes("Hanui-Relay-4min.pptx", "Hanui-Relay-4min-script.md")
    add_speaker_notes("Hanui-Relay-10min.pptx", "Hanui-Relay-10min-script.md")
    for temporary_crop in ASSETS.glob("*-crop-*.png"):
        temporary_crop.unlink()
